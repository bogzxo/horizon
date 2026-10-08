#!/usr/bin/env python3
"""Builds the OpenGL against Vulkan workbook out of the JSON run_scenes_windows.py wrote for each build.

usage: make_vulkan_workbook.py <metrics-dir> <out.xlsx>
"""
import json, os, statistics, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

metrics, out = sys.argv[1], sys.argv[2]

def load(name):
    p = os.path.join(metrics, name)
    return json.load(open(p)) if os.path.exists(p) else None

gl = load("gl_windows.json")
vk = load("vulkan_windows.json")
if gl is None or vk is None:
    sys.exit("both gl_windows.json and vulkan_windows.json are wanted")

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
GOOD = PatternFill("solid", fgColor="D9EAD3")
BAD = PatternFill("solid", fgColor="F4CCCC")
NOTE = Font(italic=True, color="555555")
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def sheet(title, headers, rows, widths=None, notes=None, colour_col=None, higher_is_better=True):
    ws = wb.create_sheet(title)
    r = 1
    for n in notes or []:
        ws.cell(row=r, column=1, value=n).font = NOTE
        r += 1
    if notes: r += 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEAD; cell.fill = HEAD_FILL; cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BORDER
    ws.row_dimensions[r].height = 32
    head = r
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if isinstance(v, float):
                cell.number_format = "0.00"
        if colour_col is not None and isinstance(row[colour_col], (int, float)):
            d = row[colour_col]
            better = d > 0 if higher_is_better else d < 0
            if abs(d) >= 1.0:
                ws.cell(row=r, column=colour_col + 1).fill = GOOD if better else BAD
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths[c - 1] if widths else 16)
    ws.freeze_panes = ws.cell(row=head + 1, column=2)
    return ws

def median(runs, loop, key):
    vals = [run[loop][key] for run in runs if run.get(loop)]
    return statistics.median(vals) if vals else None

def setup(runs):
    vals = [r["setup_ms"] for r in runs if r.get("setup_ms") is not None]
    return statistics.median(vals) if vals else None

summary = wb.active
summary.title = "Summary"

rows = []
speedups = []
for scene in vk["scenes"]:
    g, v = gl["scenes"].get(scene, []), vk["scenes"].get(scene, [])
    if not v: continue
    g_fps, v_fps = median(g, "render", "rate"), median(v, "render", "rate")
    g_ms, v_ms = median(g, "render", "ms"), median(v, "render", "ms")
    g_worst, v_worst = median(g, "render", "worst_ms"), median(v, "render", "worst_ms")
    g_alloc, v_alloc = median(g, "render", "bytes_per_turn"), median(v, "render", "bytes_per_turn")
    g_logic, v_logic = median(g, "logic", "ms"), median(v, "logic", "ms")
    g_setup, v_setup = setup(g), setup(v)
    g_checks = "; ".join(f"{a}/{t}" for r in g for a, t in r["checks"]) or ""
    v_checks = "; ".join(f"{a}/{t}" for r in v for a, t in r["checks"]) or ""
    g_errors = sum(len(r["errors"]) for r in g); v_errors = sum(len(r["errors"]) for r in v)
    speedup = v_fps / g_fps if g_fps and v_fps else None
    ms_delta = (v_ms - g_ms) / g_ms * 100.0 if g_ms and v_ms else None
    if speedup: speedups.append(speedup)
    rows.append([scene, g_fps, v_fps, speedup, g_ms, v_ms, ms_delta, g_worst, v_worst, g_logic, v_logic,
                 g_alloc, v_alloc, g_setup, v_setup, g_errors, v_errors, g_checks, v_checks])

sheet("Scenes",
      ["Scene", "FPS OpenGL", "FPS Vulkan", "Speed up (x)", "Render ms/frame OpenGL", "Render ms/frame Vulkan", "Render ms delta %",
       "Worst frame ms OpenGL", "Worst frame ms Vulkan", "Logic ms/tick OpenGL", "Logic ms/tick Vulkan",
       "Render bytes/frame OpenGL", "Render bytes/frame Vulkan", "Scene set-up ms OpenGL", "Scene set-up ms Vulkan",
       "Errors logged OpenGL", "Errors logged Vulkan", "Self-checks OpenGL", "Self-checks Vulkan"],
      rows,
      widths=[22, 11, 11, 11, 13, 13, 12, 12, 12, 12, 12, 14, 14, 12, 12, 10, 10, 14, 14],
      notes=[f"Every example scene of Horizon.Testing run for {vk['seconds']} seconds on each build, the steady half of the engine's own loop log (HORIZON_LOG_LOOPS) averaged, the median over passes.",
             "OpenGL is project-health-2 (9506ed9, OpenGL 4.6 through Silk.NET), Vulkan is the vulkan branch (Vulkan 1.3 through Silk.NET, Slang shaders).",
             "Windows 11, NVIDIA GeForce RTX 3060 (driver 610.88), Ryzen 5 5600X, a 1600 by 900 window, vsync off, no frame limit, shader cache off so set-up times include compiling.",
             "Render ms/frame is the render thread's smoothed work time per frame as the engine logs it, with the CPU side of the frame and whatever it waited for the GPU in it.",
             "FPS is frames the render loop managed a second. Scenes that stay under a few hundred frames a second are GPU bound on one or the other, the rest are CPU bound and say how much the driver and the engine cost a frame.",
             "Bytes/frame is the engine's smoothed allocation figure for the render thread. The simulation ran at its 120 Hz target in every run on both builds.",
             "The lighting, quickstart and pathtraced scenes do not draw the same thing on the two builds. Vulkan has the radiance cascade tracer, soft shadows and the sprite shadow field, OpenGL has the old path tracer and hard shadows."],
      colour_col=3)

# ---------------------------------------------------------------- Summary
summary.cell(row=1, column=1, value="Horizon, OpenGL against Vulkan").font = Font(bold=True, size=14)
summary.cell(row=2, column=1, value="The same example scenes on the same machine, the old GL 4.6 backend (project-health-2) against the Vulkan one (vulkan branch).").font = NOTE
summary.cell(row=3, column=1, value="See the Scenes sheet for every scene and the Method sheet for how the numbers were taken.").font = NOTE

lines = []
if speedups:
    lines.append(("Scenes measured on both builds", len(speedups)))
    lines.append(("Median speed up (frames a second, Vulkan over OpenGL)", f"{statistics.median(speedups):.2f}x"))
    lines.append(("Geometric mean speed up", f"{statistics.geometric_mean(speedups):.2f}x"))
    lines.append(("Scenes faster on Vulkan", sum(1 for s in speedups if s > 1.02)))
    lines.append(("Scenes slower on Vulkan", sum(1 for s in speedups if s < 0.98)))
    best = max(rows, key=lambda r: r[3] or 0); worst = min(rows, key=lambda r: r[3] or 9)
    lines.append(("Biggest gain", f"{best[0]} ({best[3]:.2f}x, {best[1]:.0f} to {best[2]:.0f} fps)"))
    lines.append(("Smallest gain", f"{worst[0]} ({worst[3]:.2f}x, {worst[1]:.0f} to {worst[2]:.0f} fps)"))
    g_allocs = [r[11] for r in rows if r[11] is not None]; v_allocs = [r[12] for r in rows if r[12] is not None]
    if g_allocs and v_allocs:
        lines.append(("Median render thread allocation a frame", f"OpenGL {statistics.median(g_allocs):.0f} bytes, Vulkan {statistics.median(v_allocs):.0f} bytes"))
    g_setups = [r[13] for r in rows if r[13] is not None]; v_setups = [r[14] for r in rows if r[14] is not None]
    if g_setups and v_setups:
        lines.append(("Median scene set-up (shaders compiled from scratch)", f"OpenGL {statistics.median(g_setups):.0f} ms, Vulkan {statistics.median(v_setups):.0f} ms"))

r = 5
for name, value in lines:
    summary.cell(row=r, column=1, value=name).font = Font(bold=True)
    summary.cell(row=r, column=2, value=value)
    r += 1

r += 1
summary.cell(row=r, column=1, value="Scene").font = HEAD; summary.cell(row=r, column=1).fill = HEAD_FILL
summary.cell(row=r, column=2, value="FPS OpenGL").font = HEAD; summary.cell(row=r, column=2).fill = HEAD_FILL
summary.cell(row=r, column=3, value="FPS Vulkan").font = HEAD; summary.cell(row=r, column=3).fill = HEAD_FILL
summary.cell(row=r, column=4, value="Speed up").font = HEAD; summary.cell(row=r, column=4).fill = HEAD_FILL
for row in sorted(rows, key=lambda x: -(x[3] or 0)):
    r += 1
    summary.cell(row=r, column=1, value=row[0])
    summary.cell(row=r, column=2, value=row[1]).number_format = "0"
    summary.cell(row=r, column=3, value=row[2]).number_format = "0"
    c = summary.cell(row=r, column=4, value=row[3]); c.number_format = "0.00"
    if row[3]: c.fill = GOOD if row[3] > 1.02 else BAD if row[3] < 0.98 else PatternFill()
summary.column_dimensions["A"].width = 56
summary.column_dimensions["B"].width = 44
summary.column_dimensions["C"].width = 12
summary.column_dimensions["D"].width = 10

# ---------------------------------------------------------------- Method
method = wb.create_sheet("Method")
for i, line in enumerate([
    "How the numbers were taken",
    "",
    "Two builds of Horizon.Testing, the example scenes of the engine, on the same machine on the same day.",
    "OpenGL is the project-health-2 branch at 9506ed9, the last state of the OpenGL 4.6 backend, built from a worktree of it.",
    "Vulkan is the vulkan branch, the engine rebuilt around Vulkan 1.3 with Slang shaders, built the same way.",
    "",
    "Machine. Windows 11 Pro, NVIDIA GeForce RTX 3060 on driver 610.88, AMD Ryzen 5 5600X. A 1600 by 900 window, vsync off, no frame limit.",
    "",
    f"Each scene was started with HORIZON_LOG_LOOPS=1, which makes the engine log its loops once a second, and HORIZON_INPUT_SCRIPT quitting it after {vk['seconds']} seconds.",
    "HORIZON_SHADER_CACHE=off, so every run compiles its shaders, which is where the set-up time goes on both builds.",
    "The second half of the log lines of a run are averaged, the first half has the scene being set up and the shaders compiling in it.",
    "Where a scene was run more than once the median of the runs is shown.",
    "",
    "The render figures are the render thread's own accounting, its work time per frame (smoothed) and frames a second, which has whatever the thread waited for the GPU in it.",
    "The Vulkan build also has GPU timestamps per pass, which the performance overlay shows (F3, Full) but which this run did not collect.",
    "",
    "What is not the same. The lighting, quickstart and pathtraced scenes draw differently on the two builds, Vulkan does radiance cascades, soft shadows and the sprite shadow field, OpenGL the old tracer and hard shadows.",
    "The CPU microbenchmarks of the project-health-2 workbook (sprite blending, colour lerps, the entity tree, HIDL) are not run again, nothing in them touches the GPU.",
    "",
    "The scripts are in docs/metrics/raw, run_scenes_windows.py took the numbers and make_vulkan_workbook.py made this file.",
], start=1):
    cell = method.cell(row=i, column=1, value=line)
    if i == 1: cell.font = Font(bold=True, size=13)
method.column_dimensions["A"].width = 160

wb.save(out)
print(f"wrote {out}")
