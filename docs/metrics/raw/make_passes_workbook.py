#!/usr/bin/env python3
"""Builds the lighting performance workbook, what every GPU pass cost before and after, out of the JSON
run_passes.py wrote, plus the frame rates run_scenes_windows.py measured on bogz's machine for the scenes that
started it (post and lighting, OpenGL against Vulkan) so the two sit side by side.

usage: make_passes_workbook.py <metrics-raw-dir> <out.xlsx>
"""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

raw, out = sys.argv[1], sys.argv[2]

def load(name):
    p = os.path.join(raw, name)
    return json.load(open(p)) if os.path.exists(p) else None

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
GOOD = PatternFill("solid", fgColor="D9EAD3")
BAD = PatternFill("solid", fgColor="F4CCCC")
NOTE = Font(italic=True, color="555555")
BOLD = Font(bold=True)
thin = Side(style="thin", color="BBBBBB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

def sheet(title, headers, rows, widths=None, notes=None, colour_col=None, lower_is_better=True):
    ws = wb.create_sheet(title)
    r = 1
    for n in notes or []:
        ws.cell(row=r, column=1, value=n).font = NOTE
        r += 1
    if notes:
        r += 1
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=r, column=c, value=h)
        cell.font = HEAD; cell.fill = HEAD_FILL; cell.alignment = Alignment(wrap_text=True, vertical="center"); cell.border = BORDER
    ws.row_dimensions[r].height = 32
    head = r
    for row in rows:
        r += 1
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = BORDER
            if isinstance(v, float):
                cell.number_format = "0.0"
        if colour_col is not None and isinstance(row[colour_col], (int, float)):
            d = row[colour_col]
            better = d < 0 if lower_is_better else d > 0
            if abs(d) >= 1.0:
                ws.cell(row=r, column=colour_col + 1).fill = GOOD if better else BAD
    for c in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths[c - 1] if widths else 16)
    ws.freeze_panes = ws.cell(row=head + 1, column=2)
    return ws

def pct(a, b):
    return None if not a or b is None else (b - a) / a * 100.0

before = load("passes_before.json")
after_examples = load("passes_after_examples.json")
japan = {
    "pt": load("passes_after_japan_pt.json"),
    "pt_nocells": load("passes_after_japan_pt_nocells.json"),
    "oldmap_pt": load("passes_after_japan_oldmap_pt.json"),
    "direct": load("passes_after_japan_direct.json"),
    "direct_nocells": load("passes_after_japan_direct_nocells.json"),
    "oldmap_direct": load("passes_after_japan_oldmap_direct.json"),
}
gl = load("gl_windows.json")
vk = load("vulkan_windows.json")

def run(data, name):
    if data is None: return None
    runs = data["runs"]
    if name in runs: return runs[name]
    for k, v in runs.items():
        if k.startswith(name): return v
    return None

# 1. The story
ws = wb.active
ws.title = "Read me"
lines = [
    "Lighting performance, before and after, October 2026 (project-health-2 after the Vulkan merge).",
    "",
    "What started it. On bogz's machine the post scene went from 953 fps on OpenGL to 244 on Vulkan and the lighting scene from 820 to 201,",
    "while every CPU bound scene got about twice as fast. The GPU passes (HORIZON_LOG_LOOPS prints them now) said why on the first run.",
    "",
    "What was wrong. The field of sprite shadows (a signed distance field of every sprite that blocks light, built every frame) was three",
    "passes of which two were serial, one thread per column sweeping the whole column and one thread per row building an envelope, on",
    "one lane of a workgroup. That is the one shape of work a GPU can't do, and it ran on every frame of every deferred renderer whether",
    "or not anything cast a shadow or any light was in view, the post scene has no lights at all. The path tracer's rays each worked out",
    "the lamps with their shadows at the wall they hit, the whole deferred pass over again per ray.",
    "",
    "What changed. The field is a jump flood now (sprite_seed, sprite_flood eight times, sprite_field), on a grid two texels coarse with the",
    "last pass exact, and it is only built when a sprite that casts shadows was drawn and a light is in view (or the tracer is on, it",
    "wants the field for what glows). The tracer works out what every wall throws back once a frame (gi_radiance) and the rays read it.",
    "The light tiles are skipped with no lights. And with a LightingPixelSize set (pixel art, Fighter2D lights a pixel of the art at a",
    "time) the lighting is worked out once per pixel of the art by a compute pass (light_cells) and the deferred pass reads it, instead",
    "of every pixel of the screen marching the same shadows to the same lights, which at the fight's zoom is seven pixels for one.",
    "",
    "How it was measured. Headless on lavapipe (Mesa's CPU Vulkan) at 1600 by 900 on a four core box, which is slow but times every pass",
    "honestly, the shares and the ratios are what matter. 'Lavapipe passes' is the engine examples before and after on the same scenes.",
    "'Japan' is the game on its map, path traced and direct, with the new map (raised lanterns as swaying spots and a moon) and the old",
    "one (the old lanterns as point lights, no moon), and with the per art pixel lighting on and off, so the engine's share and the map's",
    "share can be told apart. 'Windows before' is what bogz measured on his GPU before any of this, OpenGL against Vulkan, kept here for",
    "the reference. The after on that machine takes a run of docs/metrics/raw/run_scenes_windows.py there, which is the real test.",
    "",
    "Not measured here. The tile map and the UI passes are what they were, the resolve of a scene without a LightingPixelSize is what",
    "it was (the lighting example's resolve is the same code as before), and nothing in here says what a real GPU makes of the jump",
    "flood's eight passes, they are a few hundred thousand cheap threads each where the old passes were a few thousand long ones.",
]
for i, l in enumerate(lines, 1):
    ws.cell(row=i, column=1, value=l).font = BOLD if i == 1 else Font()
ws.column_dimensions["A"].width = 140

# 2. Examples before and after
rows = []
if before and after_examples:
    for scene in ["post", "lighting", "pathtraced"]:
        b, a = run(before, scene), run(after_examples, scene)
        if not b or not a: continue
        names = []
        for n in list(b["passes"].keys()) + list(a["passes"].keys()):
            if n not in names: names.append(n)
        rows.append([scene, "whole frame", b["gpu_ms"], a["gpu_ms"], pct(b["gpu_ms"], a["gpu_ms"])])
        for n in names:
            bv, av = b["passes"].get(n), a["passes"].get(n)
            rows.append(["", n.replace(">", "   "), bv, av if av is not None else 0.0, pct(bv, av if av is not None else 0.0) if bv else None])
sheet("Lavapipe passes", ["Scene", "Pass (indented is inside the one above)", "Before, ms", "After, ms", "Change, %"], rows,
      widths=[14, 44, 14, 14, 12], colour_col=4,
      notes=["The engine's example scenes headless on lavapipe at 1600 by 900, what the GPU spent on every pass of a frame, averaged over the steady half of a 10 to 12 second run.",
             "A pass at 0 after is one that no longer runs for that scene (the post scene has no lights, so no sprite shadow field and no light tiles)."])

# 3. Japan
rows = []
sets = [
    ("new map, path traced", japan["pt"]), ("new map, path traced, cells off", japan["pt_nocells"]), ("old map, path traced", japan["oldmap_pt"]),
    ("new map, direct", japan["direct"]), ("new map, direct, cells off", japan["direct_nocells"]), ("old map, direct", japan["oldmap_direct"]),
]
b = run(before, "--map japan")
if b:
    rows.append(["before, old map, path traced, old engine", "whole frame", b["gpu_ms"]])
    for n, v in b["passes"].items():
        rows.append(["", n.replace(">", "   "), v])
for label, data in sets:
    r = run(data, "--map") if data else None
    if not r: continue
    rows.append([label, "whole frame", r["gpu_ms"]])
    for n, v in r["passes"].items():
        if n.startswith(">particles") and v < 1.0: continue
        rows.append(["", n.replace(">", "   "), v])
sheet("Japan", ["Run", "Pass", "ms"], rows, widths=[44, 40, 12],
      notes=["Fighter2D on the Japan map headless on lavapipe, the versus screen and the first seconds of the fight, 12 seconds a run.",
             "'before' is the engine as merged with the old map. 'old map' is the new engine on that same map, which is the engine's doing alone.",
             "'new map' has the lanterns raised and turned into swaying spots, a moon (a directional light, long shadows over everything) and a dusk ambient.",
             "'cells off' (HORIZON_LIGHT_CELLS=off) has the deferred pass light every pixel itself, the way it did before, on the same map, so the per art pixel lighting can be weighed on its own."])

# 4. Windows before, for the record
rows = []
if gl and vk:
    for scene in ["post", "lighting", "pathtraced", "entities", "sprites", "tilemap", "particles", "ui-screens", "fluid"]:
        g, v = gl["scenes"].get(scene), vk["scenes"].get(scene)
        if not g or not v: continue
        gr = [p["render"]["rate"] for p in g if p.get("render")]
        vr = [p["render"]["rate"] for p in v if p.get("render")]
        if not gr or not vr: continue
        gf, vf = sum(gr) / len(gr), sum(vr) / len(vr)
        rows.append([scene, gf, vf, pct(gf, vf)])
sheet("Windows before", ["Scene", "OpenGL, fps", "Vulkan, fps (before this work)", "Change, %"], rows, widths=[16, 16, 28, 12], colour_col=3, lower_is_better=False,
      notes=["What bogz measured on his machine before any of this (gl_windows.json and vulkan_windows.json), the frame rates of the example scenes.",
             "The post and lighting rows are the ones that started it. Run run_scenes_windows.py on that machine on the new build for the after."])

wb.save(out)
print(f"wrote {out}")
